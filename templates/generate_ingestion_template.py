"""Generate the GDC application ingestion Excel template with a Jira example."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

OUTPUT = Path(__file__).parent / "gdc-app-ingestion-template.xlsx"
MAX_ROWS = 200

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
SECTION_FILL = PatternFill("solid", fgColor="D9E1F2")
EXAMPLE_FILL = PatternFill("solid", fgColor="FFF2CC")
REQUIRED_FILL = PatternFill("solid", fgColor="FCE4D6")
HEADER_FONT = Font(bold=True, color="FFFFFF")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")

LOOKUPS: dict[str, list[str]] = {
    "YesNo": ["Yes", "No", "N/A", "TBD"],
    "AppType": ["COTS (Vendor)", "Open Source", "In-house", "SaaS Integration"],
    "DeployTarget": [
        "K8s - Helm Chart",
        "K8s - Docker Image (raw manifests)",
        "K8s - Kustomize",
        "VM - Ansible",
        "Terraform - Infra Provisioning",
        "Terraform - Day-2 Configuration",
        "Platform Shared Service",
    ],
    "ProvidedBy": ["Vendor", "Open Source Community", "Platform Team", "App Team"],
    "ArtifactType": [
        "Helm Chart",
        "Docker/OCI Image",
        "Debian Package",
        "RPM Package",
        "Ansible Role",
        "Ansible Collection",
        "Terraform Provider",
        "Terraform Module",
        "VM / OS Image",
        "Generic Binary / Archive",
        "GPG Key / Certificate",
        "Python Package",
        "Maven / JAR",
        "npm Package",
    ],
    "JFrogPkgType": [
        "helm",
        "docker",
        "oci",
        "debian",
        "rpm",
        "ansible",
        "terraform",
        "generic",
        "pypi",
        "maven",
        "npm",
    ],
    "JFrogRepoType": ["Remote", "Local", "Virtual", "Federated"],
    "IngestionMethod": [
        "Remote repo proxy/cache",
        "Pull-Scan-Push to Local",
        "Manual upload (vendor portal)",
        "Built internally (CI publish)",
        "Transfer via data diode / air-gap media",
    ],
    "Environment": ["Dev", "Test", "Stage", "Prod", "Dev, Test, Prod", "All"],
    "Criticality": ["Tier 0 - Mission Critical", "Tier 1 - Business Critical", "Tier 2 - Important", "Tier 3 - Low"],
    "DataClass": ["Public", "Internal", "Confidential", "Restricted"],
    "TFPurpose": ["Infra Provisioning", "Day-2 Configuration", "Platform Bootstrap"],
    "SecretStore": ["HashiCorp Vault", "K8s Secret (ExternalSecrets)", "Ansible Vault", "GitLab CI Variable (masked)", "JFrog Access Token"],
    "Status": ["Not Started", "In Progress", "Done", "Blocked", "N/A"],
    "Direction": ["Inbound", "Outbound", "Internal (east-west)"],
}


def style_header(ws: Worksheet, row: int, ncols: int) -> None:
    for col in range(1, ncols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = BORDER


def set_widths(ws: Worksheet, widths: list[int]) -> None:
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width


def add_list_validation(ws: Worksheet, lookup: str, col_letter: str, first_row: int) -> None:
    values = LOOKUPS[lookup]
    ref = f"=Lookups!${lookup_col(lookup)}$2:${lookup_col(lookup)}${len(values) + 1}"
    dv = DataValidation(type="list", formula1=ref, allow_blank=True, showErrorMessage=False)
    dv.add(f"{col_letter}{first_row}:{col_letter}{MAX_ROWS}")
    ws.add_data_validation(dv)


def lookup_col(lookup: str) -> str:
    return get_column_letter(list(LOOKUPS).index(lookup) + 1)


def build_table_sheet(
    wb: Workbook,
    title: str,
    headers: list[tuple[str, int, bool]],
    rows: list[list[str]],
    validations: dict[str, str] | None = None,
) -> Worksheet:
    """Row-oriented sheet. headers = (name, width, required). Example rows are highlighted."""
    ws = wb.create_sheet(title)
    ws.append([f"{h}{' *' if req else ''}" for h, _, req in headers])
    style_header(ws, 1, len(headers))
    for row in rows:
        ws.append(row)
        r = ws.max_row
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=r, column=col)
            cell.fill = EXAMPLE_FILL
            cell.alignment = WRAP
            cell.border = BORDER
    set_widths(ws, [w for _, w, _ in headers])
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{MAX_ROWS}"
    ws.row_dimensions[1].height = 45
    for col_letter, lookup in (validations or {}).items():
        add_list_validation(ws, lookup, col_letter, 2)
    return ws


def build_form_sheet(
    wb: Workbook,
    title: str,
    value_headers: list[str],
    sections: list[tuple[str, list[list[str]]]],
    widths: list[int] | None = None,
) -> Worksheet:
    """Attribute-oriented sheet: Attribute | Description | Required | <value columns>."""
    ws = wb.create_sheet(title)
    headers = ["Attribute", "Description / Guidance", "Required"] + value_headers
    ws.append(headers)
    style_header(ws, 1, len(headers))
    for section, items in sections:
        ws.append([section])
        r = ws.max_row
        for col in range(1, len(headers) + 1):
            ws.cell(row=r, column=col).fill = SECTION_FILL
            ws.cell(row=r, column=col).font = Font(bold=True)
        for item in items:
            ws.append(item)
            r = ws.max_row
            for col in range(1, len(headers) + 1):
                cell = ws.cell(row=r, column=col)
                cell.alignment = WRAP
                cell.border = BORDER
                if col > 3:
                    cell.fill = EXAMPLE_FILL
            if len(item) > 2 and item[2] == "Yes":
                ws.cell(row=r, column=3).fill = REQUIRED_FILL
    set_widths(ws, widths or [38, 55, 10] + [55] * len(value_headers))
    ws.freeze_panes = "D2"
    return ws


def build_instructions(wb: Workbook) -> None:
    ws = wb.active
    ws.title = "0_Instructions"
    lines = [
        ("GDC Application Ingestion Template (JFrog-backed Self-Service)", True),
        ("Purpose: capture everything the platform team needs to ingest an application's artifacts into JFrog "
         "and generate the deployment pipeline (K8s / VM / Terraform) onto GDC.", False),
        ("", False),
        ("How to fill in", True),
        ("1. Start with 1_App_Overview (one per application).", False),
        ("2. List every deployable piece in 2_Components (K8s workloads, VMs, Terraform stacks, day-2 config).", False),
        ("3. List EVERY artifact that must exist in JFrog in 3_Artifact_Inventory, including transitive ones "
         "(e.g. images referenced inside a Helm chart, init/sidecar images, OS packages, Terraform providers, Ansible collections).", False),
        ("4. Complete 4_K8s_Deployment for each K8s component and 5_VM_Deployment for each VM component.", False),
        ("5. Complete 6_Terraform for every Terraform stack (infra provisioning AND day-2 config such as GitLab/Jira projects, groups).", False),
        ("6. Complete 7_Network_Flows, 8_Secrets_Config, 9_Deploy_Sequence and 10_Readiness_Checklist.", False),
        ("", False),
        ("Conventions", True),
        ("- Columns/attributes marked * or Required=Yes are mandatory for ingestion approval.", False),
        ("- Yellow cells are the worked EXAMPLE (Jira Software Data Center). Replace or append your own rows.", False),
        ("- Values marked 'confirm' are illustrative and must be verified against vendor docs / platform standards.", False),
        ("- Never put secret VALUES in this workbook; reference the secret store path only.", False),
        ("- Pin every artifact to an exact version, and to a digest/checksum where possible (no 'latest').", False),
        ("- Use dropdowns where provided; lists are maintained in the Lookups sheet.", False),
        ("", False),
        ("Example scenario: Jira", True),
        ("Jira Software DC runs on GDC K8s using the vendor (Atlassian) Helm chart. GDC has no managed PostgreSQL 17, "
         "so PostgreSQL 17 is deployed on a GDC VM, provisioned by Terraform and configured by Ansible using Debian packages "
         "proxied through JFrog. Jira projects/groups are created day-2 via a Terraform provider.", False),
    ]
    for text, bold in lines:
        ws.append([text])
        cell = ws.cell(row=ws.max_row, column=1)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if bold:
            cell.font = Font(bold=True, size=12, color="1F3864")
    ws.column_dimensions["A"].width = 140


def build_overview(wb: Workbook) -> None:
    sections = [
        ("Identification", [
            ["Ingestion Request ID", "Unique ID (e.g. Jira ticket key of the request)", "Yes", "ING-2026-001"],
            ["Application Name", "Full product name", "Yes", "Jira Software Data Center"],
            ["Application Short Code", "Lowercase, used for namespaces, repo keys, DNS", "Yes", "jira"],
            ["Business Description", "What the app does / who uses it", "Yes", "Issue and project tracking for engineering teams"],
            ["Application Type", "COTS / OSS / In-house", "Yes", "COTS (Vendor)"],
            ["Vendor / Upstream Source", "Vendor or project name", "Yes", "Atlassian"],
            ["Application Version", "Exact product version to ingest", "Yes", "10.3.x LTS (example - confirm exact patch)"],
            ["Vendor Documentation URL", "Install / Helm / support matrix docs", "Yes", "https://atlassian.github.io/data-center-helm-charts/"],
            ["Deployment Pattern", "K8s only / VM only / Hybrid", "Yes", "Hybrid: K8s (Helm) + VM (PostgreSQL) + Terraform"],
        ]),
        ("Ownership & Support", [
            ["Business Owner", "Name / email", "Yes", "<name@org>"],
            ["Technical Owner", "Name / email", "Yes", "<name@org>"],
            ["Support Team / On-call", "Team / rotation / escalation", "Yes", "Collaboration Tools Team"],
            ["Requesting Team GitLab Group", "Where app config/IaC repos live", "Yes", "gitlab/collab-tools/jira"],
            ["Cost Center", "", "No", "<cost-center>"],
            ["Vendor Support Contract", "Contract ID / expiry", "No", "<contract-id>"],
        ]),
        ("Classification & Compliance", [
            ["Data Classification", "Public / Internal / Confidential / Restricted", "Yes", "Confidential"],
            ["Criticality Tier", "Drives HA / DR / change process", "Yes", "Tier 1 - Business Critical"],
            ["Compliance Requirements", "e.g. FedRAMP, ISO 27001, CIS, STIG", "Yes", "CIS hardened OS, audit logging enabled"],
            ["License Type", "Commercial / OSS licence (Apache-2.0, GPL ...)", "Yes", "Commercial (Atlassian DC subscription)"],
            ["License Key Handling", "Where licence key is stored", "Yes", "Vault: secret/apps/jira/license"],
        ]),
        ("Target Environment", [
            ["Target Environments", "Dev / Test / Stage / Prod", "Yes", "Dev, Test, Prod"],
            ["GDC Organization / Project", "GDC org & project identifiers", "Yes", "org-collab / prj-jira"],
            ["GDC Zone(s)", "Zones for deployment", "Yes", "zone-1 (Prod: zone-1 + zone-2)"],
            ["K8s Cluster(s)", "User cluster name(s)", "Yes", "user-cluster-apps-01"],
            ["Air-gapped?", "Is upstream internet reachable only via JFrog?", "Yes", "Yes - all artifacts via JFrog only"],
            ["Expected Users / Load", "Concurrent users, growth", "Yes", "2,000 named / 400 concurrent"],
            ["HA Requirement", "Replicas, multi-zone", "Yes", "Jira 2+ pods (DC clustering), PG primary + standby"],
            ["RTO / RPO", "", "Yes", "RTO 4h / RPO 15m"],
            ["Target Go-Live Date", "", "No", "2026-12-15"],
        ]),
    ]
    build_form_sheet(wb, "1_App_Overview", ["Value (Example: Jira)"], sections)


def build_components(wb: Workbook) -> None:
    headers = [
        ("Component ID", 12, True), ("Component Name", 28, True), ("Description", 45, True),
        ("Deployment Target", 28, True), ("Provided By", 18, True), ("Depends On (Component IDs)", 18, False),
        ("Environments", 16, True), ("Why this target? (justification)", 45, False), ("Owner", 18, True), ("Notes", 40, False),
    ]
    rows = [
        ["C01", "Jira Software DC", "Jira application pods (StatefulSet)", "K8s - Helm Chart", "Vendor", "C02, C04, C05, C07",
         "Dev, Test, Prod", "Vendor ships official Helm chart", "Collab Tools Team", "Chart: atlassian-data-center/jira"],
        ["C02", "PostgreSQL 17 database", "Jira backend database", "VM - Ansible", "Open Source Community", "C03",
         "Dev, Test, Prod", "GDC has no managed PostgreSQL 17 service", "Platform DBA", "Packages from apt.postgresql.org via JFrog Debian remote"],
        ["C03", "PostgreSQL VM infrastructure", "VM(s), disks, network for PostgreSQL", "Terraform - Infra Provisioning", "Platform Team", "",
         "Dev, Test, Prod", "Standard platform VM module", "Platform Team", "Primary + standby in Prod"],
        ["C04", "Jira K8s namespace baseline", "Namespace, quotas, limit ranges, network policies, imagePullSecret", "Terraform - Infra Provisioning",
         "Platform Team", "", "Dev, Test, Prod", "Self-service namespace onboarding", "Platform Team", ""],
        ["C05", "Jira shared home storage", "RWX volume for Jira DC shared home", "Platform Shared Service", "Platform Team", "C04",
         "Dev, Test, Prod", "Required by Jira DC clustering", "Platform Team", "RWX StorageClass (NFS/file) - confirm availability on GDC"],
        ["C06", "Jira day-2 configuration", "Projects, groups, memberships, permission schemes", "Terraform - Day-2 Configuration",
         "App Team", "C01", "Dev, Test, Prod", "Config-as-code instead of manual admin", "Collab Tools Team", "Uses Jira Terraform provider (community - security review)"],
        ["C07", "Jira ingress & TLS", "Ingress host, TLS certificate", "Platform Shared Service", "Platform Team", "C04",
         "Dev, Test, Prod", "Standard ingress", "Platform Team", "Cert from internal CA / cert-manager"],
        ["C08", "GitLab project for Jira IaC", "GitLab group/project holding Jira Helm values & Terraform", "Terraform - Day-2 Configuration",
         "Platform Team", "", "All", "Self-service repo bootstrap", "Platform Team", "gitlabhq/gitlab provider"],
    ]
    build_table_sheet(wb, "2_Components", headers, rows,
                      {"D": "DeployTarget", "E": "ProvidedBy", "G": "Environment"})


def build_artifacts(wb: Workbook) -> None:
    headers = [
        ("Artifact ID", 10, True), ("Component ID", 11, True), ("Artifact Type", 20, True), ("Artifact Name", 30, True),
        ("Version / Tag", 18, True), ("Upstream Source (registry / URL)", 38, True), ("Full Upstream Reference", 45, True),
        ("Digest / SHA256 Checksum", 28, True), ("Signature / Provenance Available", 16, False), ("SBOM Available", 12, False),
        ("License", 18, True), ("JFrog Package Type", 14, True), ("JFrog Repo Type", 12, True), ("Target JFrog Repo Key", 30, True),
        ("Ingestion Method", 26, True), ("Reference After Ingestion (JFrog path)", 50, True), ("Xray Scan Required", 12, True),
        ("Referenced By (transitive)", 22, False), ("Optional?", 10, False), ("Notes", 45, False),
    ]
    rows = [
        ["A01", "C01", "Helm Chart", "jira", "1.x (example - confirm)", "https://atlassian.github.io/data-center-helm-charts",
         "atlassian-data-center/jira", "<sha256 of .tgz>", "No", "No", "Apache-2.0", "helm", "Remote",
         "helm-remote-atlassian -> helm-virtual", "Remote repo proxy/cache", "https://<jfrog>/artifactory/api/helm/helm-virtual (chart: jira)",
         "Yes", "", "No", "Override image.repository to JFrog docker-virtual in values"],
        ["A02", "C01", "Docker/OCI Image", "atlassian/jira-software", "10.3.x (example - confirm)", "docker.io",
         "docker.io/atlassian/jira-software:<tag>", "sha256:<digest>", "TBD", "TBD", "Commercial", "docker", "Remote",
         "docker-remote-dockerhub -> docker-virtual", "Pull-Scan-Push to Local", "<jfrog>/docker-virtual/atlassian/jira-software:<tag>",
         "Yes", "A01 (image.repository)", "No", "Pin by digest in values"],
        ["A03", "C01", "Docker/OCI Image", "alpine (nfs permission fixer init container)", "<pinned tag>", "docker.io",
         "docker.io/library/alpine:<tag>", "sha256:<digest>", "No", "No", "MIT", "docker", "Remote",
         "docker-remote-dockerhub -> docker-virtual", "Pull-Scan-Push to Local", "<jfrog>/docker-virtual/library/alpine:<tag>",
         "Yes", "A01 (volumes.sharedHome.nfsPermissionFixer)", "Yes", "Confirm via 'helm template' whether enabled"],
        ["A04", "C01", "Docker/OCI Image", "fluentd sidecar", "<pinned tag>", "docker.io", "<per chart values: fluentd.imageRepo>",
         "sha256:<digest>", "No", "No", "Apache-2.0", "docker", "Remote", "docker-remote-dockerhub -> docker-virtual",
         "Pull-Scan-Push to Local", "<jfrog>/docker-virtual/<fluentd-image>:<tag>", "Yes", "A01 (fluentd.enabled)", "Yes",
         "Only if log shipping sidecar enabled"],
        ["A05", "C02", "Debian Package", "postgresql-17, postgresql-client-17, postgresql-common", "17.x (pinned minor)",
         "https://apt.postgresql.org/pub/repos/apt", "deb <dist>-pgdg main", "Repo Release/InRelease signed", "Yes (GPG signed repo)", "No",
         "PostgreSQL License", "debian", "Remote", "debian-remote-pgdg -> debian-virtual", "Remote repo proxy/cache",
         "https://<jfrog>/artifactory/debian-virtual <dist>-pgdg main", "Yes", "", "No", "See docs/jfrog-debian-repository.md"],
        ["A06", "C02", "GPG Key / Certificate", "PGDG repository signing key", "current", "https://www.postgresql.org/media/keys/ACCC4CF8.asc",
         "ACCC4CF8.asc", "<sha256>", "N/A", "N/A", "N/A", "generic", "Local", "generic-local/keys/pgdg/",
         "Manual upload (vendor portal)", "<jfrog>/generic-local/keys/pgdg/ACCC4CF8.asc", "No", "A05", "No", "Needed by apt on air-gapped VM"],
        ["A07", "C02", "Ansible Role", "postgresql (internal)", "v1.0.0", "Internal GitLab", "roles/postgresql (this repo)",
         "<git tag commit sha>", "N/A", "N/A", "Internal", "generic", "Local", "ansible-local", "Built internally (CI publish)",
         "<jfrog>/ansible-local/roles/postgresql-1.0.0.tar.gz", "No", "", "No", "Playbook: playbooks/postgresql.yml"],
        ["A08", "C02", "Ansible Collection", "community.postgresql", "<pinned version>", "https://galaxy.ansible.com",
         "community.postgresql:<version>", "<sha256>", "No", "No", "GPL-3.0", "ansible", "Remote", "ansible-remote-galaxy -> ansible-virtual",
         "Remote repo proxy/cache", "<jfrog>/api/ansible/ansible-virtual", "Yes", "A07", "No", "Used for DB/user creation"],
        ["A09", "C02", "VM / OS Image", "Ubuntu 22.04 LTS (CIS hardened)", "<image version>", "GDC VM image catalog", "<gdc image name>",
         "<sha256>", "TBD", "TBD", "Canonical", "generic", "Local", "vm-images-local", "Built internally (CI publish)",
         "<jfrog>/vm-images-local/ubuntu-22.04-cis/<ver>", "Yes", "", "No", "Platform golden image - confirm PG17 dist support"],
        ["A10", "C03, C04", "Terraform Provider", "hashicorp/kubernetes", "<pinned version>", "registry.terraform.io",
         "hashicorp/kubernetes", "<sha256 (lock file)>", "Yes (HashiCorp signed)", "No", "MPL-2.0", "terraform", "Remote",
         "terraform-remote-hashicorp -> terraform-virtual", "Remote repo proxy/cache", "<jfrog>/terraform-virtual (provider mirror)", "Yes", "", "No",
         "GDC VMs/namespaces via KRM API - confirm provider choice with platform"],
        ["A11", "C03", "Terraform Module", "gdc-vm (platform module)", "v2.x", "Internal GitLab", "platform/terraform-modules/gdc-vm",
         "<git tag>", "N/A", "N/A", "Internal", "terraform", "Local", "terraform-modules-local", "Built internally (CI publish)",
         "<jfrog>/terraform-modules-local/platform/gdc-vm/<ver>", "No", "", "No", ""],
        ["A12", "C06", "Terraform Provider", "Jira provider (e.g. fourplusone/jira)", "<pinned version>", "registry.terraform.io",
         "fourplusone/jira (community - confirm)", "<sha256 (lock file)>", "TBD", "No", "MPL-2.0 (confirm)", "terraform", "Remote",
         "terraform-remote-hashicorp -> terraform-virtual", "Pull-Scan-Push to Local", "<jfrog>/terraform-virtual (provider mirror)",
         "Yes", "", "No", "Community provider - security & compatibility review with Jira 10"],
        ["A13", "C08", "Terraform Provider", "gitlabhq/gitlab", "<pinned version>", "registry.terraform.io", "gitlabhq/gitlab",
         "<sha256 (lock file)>", "Yes", "No", "MPL-2.0", "terraform", "Remote", "terraform-remote-hashicorp -> terraform-virtual",
         "Remote repo proxy/cache", "<jfrog>/terraform-virtual (provider mirror)", "Yes", "", "No", "GitLab groups/projects for Jira IaC"],
    ]
    build_table_sheet(wb, "3_Artifact_Inventory", headers, rows, {
        "C": "ArtifactType", "I": "YesNo", "J": "YesNo", "L": "JFrogPkgType", "M": "JFrogRepoType",
        "O": "IngestionMethod", "Q": "YesNo", "S": "YesNo",
    })


def build_k8s(wb: Workbook) -> None:
    sections = [
        ("Packaging", [
            ["Component ID", "From 2_Components", "Yes", "C01"],
            ["Packaging Type", "Helm chart / raw manifests / Kustomize", "Yes", "Helm chart (vendor)"],
            ["Chart Name & Version", "Exact chart version", "Yes", "atlassian-data-center/jira 1.x (confirm)"],
            ["Chart Source in JFrog", "Virtual repo URL", "Yes", "helm-virtual"],
            ["Values File Location", "GitLab path per environment", "Yes", "gitlab/collab-tools/jira/helm/values-<env>.yaml"],
            ["All Images Rendered by Chart", "Output of 'helm template ... | grep image:' - every image must be in 3_Artifact_Inventory", "Yes", "A02, A03 (A04 if fluentd enabled)"],
            ["Image Registry Override Keys", "Values keys used to point images at JFrog", "Yes", "image.repository, image.tag, fluentd.imageRepo"],
            ["CRDs / Cluster-scoped Resources", "Any CRDs, ClusterRoles, webhooks", "Yes", "None"],
            ["Helm Hooks / Jobs", "Pre/post install jobs", "No", "None"],
        ]),
        ("Runtime", [
            ["Namespace", "", "Yes", "jira-<env>"],
            ["Workload Kind", "Deployment / StatefulSet / DaemonSet", "Yes", "StatefulSet"],
            ["Replicas (per env)", "", "Yes", "Dev 1 / Test 1 / Prod 2+"],
            ["CPU Request / Limit", "Per pod", "Yes", "2 / 4 (size per load - confirm)"],
            ["Memory Request / Limit", "Per pod", "Yes", "4Gi / 6Gi; JVM heap via jira.resources.jvm.maxHeap"],
            ["Container Ports", "", "Yes", "8080 (http), 40001/40011 (cluster cache)"],
            ["Service Type", "ClusterIP / LoadBalancer", "Yes", "ClusterIP"],
            ["Health Probes", "Readiness / liveness endpoints", "Yes", "/status (chart default)"],
            ["HPA / PDB", "", "No", "No HPA; PDB minAvailable 1 in Prod"],
            ["Session Affinity", "", "Yes", "Sticky sessions required when replicas > 1"],
        ]),
        ("Storage", [
            ["Local Home PVC", "Size / StorageClass / access mode", "Yes", "volumes.localHome: 20Gi, <rwo-class>, RWO"],
            ["Shared Home PVC", "Size / StorageClass / access mode", "Yes", "volumes.sharedHome: 100Gi, <rwx-class>, RWX"],
            ["Backup of PVs", "Tool / schedule", "Yes", "Platform backup (e.g. Velero/GDC backup) daily"],
        ]),
        ("Security", [
            ["Pod Security Context", "runAsUser / fsGroup / privileged?", "Yes", "fsGroup 2001, non-root, not privileged"],
            ["Service Account / RBAC", "", "Yes", "Chart-created SA, no cluster RBAC"],
            ["imagePullSecret", "JFrog pull token secret name", "Yes", "jfrog-pull (created by C04)"],
            ["Network Policies", "Ingress/egress allowed", "Yes", "Allow ingress-nginx -> 8080; egress to PG VM 5432, SMTP, LDAP"],
        ]),
        ("Integration", [
            ["Ingress Host", "", "Yes", "jira.<env>.gdc.example.org"],
            ["Ingress Class / TLS", "", "Yes", "nginx; TLS secret jira-tls (internal CA)"],
            ["Database Connection", "Values keys", "Yes", "database.type=postgres72; database.url=jdbc:postgresql://pg17-jira-<env>:5432/jiradb; database.driver=org.postgresql.Driver"],
            ["Database Credentials Secret", "", "Yes", "database.credentials.secretName=jira-db-credentials"],
            ["Other External Dependencies", "SMTP, LDAP/AD, SSO", "No", "SMTP relay, AD (LDAPS 636), SAML SSO"],
            ["Observability", "Metrics / logs / dashboards", "No", "JMX exporter (monitoring.exposeJmxMetrics), stdout to platform logging"],
        ]),
    ]
    build_form_sheet(wb, "4_K8s_Deployment", ["C01 - Jira (Example)", "<Next K8s component>"], sections)


def build_vm(wb: Workbook) -> None:
    sections = [
        ("Infrastructure", [
            ["Component ID", "From 2_Components", "Yes", "C02 (infra by C03)"],
            ["VM Role", "", "Yes", "PostgreSQL 17 database server"],
            ["Reason for VM (not K8s / managed)", "", "Yes", "GDC has no managed PostgreSQL 17; stateful DB on VM per platform standard"],
            ["OS Image & Version", "Must be in 3_Artifact_Inventory", "Yes", "Ubuntu 22.04 LTS CIS (A09) - confirm PGDG support"],
            ["VM Count per Env", "", "Yes", "Dev 1 / Test 1 / Prod 2 (primary + standby)"],
            ["vCPU / RAM", "", "Yes", "4 vCPU / 16 GiB (Prod 8 / 32)"],
            ["OS Disk", "", "Yes", "50 GiB"],
            ["Data Disk(s) & Mount", "", "Yes", "200 GiB at /var/lib/postgresql (Prod 500 GiB)"],
            ["Network / Subnet / Static IP", "", "Yes", "db-subnet; static IP; DNS pg17-jira-<env>"],
            ["Provisioned By", "Terraform stack ID from 6_Terraform", "Yes", "T01"],
        ]),
        ("Configuration (Ansible)", [
            ["Playbook", "", "Yes", "playbooks/postgresql.yml"],
            ["Roles", "", "Yes", "roles/postgresql (A07)"],
            ["Collections", "", "Yes", "community.postgresql (A08)"],
            ["Inventory Group", "", "Yes", "postgresql_jira_<env>"],
            ["Package Repos (JFrog)", "apt/yum repo URLs used on the VM", "Yes", "debian-virtual (<dist>-pgdg main) + key A06"],
            ["Packages & Versions", "", "Yes", "postgresql-17, postgresql-client-17 (pinned)"],
            ["Key Ansible Variables", "No secret values", "Yes", "postgresql_version=17; db=jiradb; owner=jirauser; encoding UNICODE; LC_COLLATE/LC_CTYPE 'C'; template0"],
            ["Service Ports", "", "Yes", "5432/tcp"],
            ["pg_hba / Allowed Clients", "", "Yes", "K8s pod/egress CIDR of user-cluster-apps-01, scram-sha-256, TLS required"],
            ["OS Service Accounts", "", "Yes", "postgres (system), ansible (automation, sudo)"],
            ["Hardening Baseline", "", "Yes", "CIS Level 1 + PostgreSQL CIS benchmark"],
        ]),
        ("Operations", [
            ["Backup Tool / Schedule / Retention", "", "Yes", "pgBackRest: full weekly, diff daily, WAL archive; 30-day retention"],
            ["HA / Replication", "", "Yes", "Streaming replication to standby (Prod)"],
            ["Monitoring", "", "Yes", "postgres_exporter + node_exporter to platform Prometheus"],
            ["Patching Window", "", "Yes", "Monthly, Sun 02:00-04:00"],
            ["Upgrade Path", "", "No", "Minor via apt pin bump; major via pg_upgrade runbook"],
            ["Decommission Playbook", "", "No", "playbooks/postgresql-uninstall.yml"],
        ]),
    ]
    build_form_sheet(wb, "5_VM_Deployment", ["C02 - PostgreSQL 17 (Example)", "<Next VM component>"], sections)


def build_terraform(wb: Workbook) -> None:
    headers = [
        ("Stack ID", 9, True), ("Component ID", 11, True), ("Purpose", 20, True), ("Description", 38, True),
        ("Code Location (GitLab repo/path)", 38, True), ("Terraform Version", 12, True), ("Providers (source@version)", 38, True),
        ("Provider Artifact IDs", 14, True), ("Modules (source@version)", 32, False), ("State Backend", 28, True),
        ("Auth / Credentials Method", 32, True), ("Resources Managed", 40, True), ("Run By (pipeline / stage)", 26, True),
        ("Approval Gate", 22, True), ("Drift Detection", 16, False), ("Environments", 14, True), ("Notes", 38, False),
    ]
    rows = [
        ["T01", "C03", "Infra Provisioning", "PostgreSQL VM(s), disks, DNS", "gitlab/collab-tools/jira/terraform/infra-vm",
         ">= 1.6 (pinned)", "hashicorp/kubernetes@<ver>", "A10", "platform/gdc-vm@v2.x (A11)", "GitLab-managed TF state (http backend)",
         "GDC service account token from Vault", "VirtualMachine, VirtualMachineDisk, DNS record", "Ingestion pipeline: provision-infra",
         "Plan review by Platform", "Nightly plan", "Dev, Test, Prod", "Confirm GDC VM API/provider with platform"],
        ["T02", "C04", "Infra Provisioning", "Namespace baseline for Jira", "gitlab/platform/namespaces/jira",
         ">= 1.6 (pinned)", "hashicorp/kubernetes@<ver>", "A10", "platform/namespace-baseline@v1.x", "GitLab-managed TF state",
         "Platform SA (Vault)", "Namespace, ResourceQuota, LimitRange, NetworkPolicy, imagePullSecret", "Ingestion pipeline: provision-infra",
         "Auto (policy-checked)", "Nightly plan", "Dev, Test, Prod", ""],
        ["T03", "C06", "Day-2 Configuration", "Jira projects, groups, memberships", "gitlab/collab-tools/jira/terraform/day2-config",
         ">= 1.6 (pinned)", "fourplusone/jira@<ver> (confirm)", "A12", "", "GitLab-managed TF state",
         "Jira admin API token from Vault", "jira_project, jira_group, jira_group_membership (confirm resources)", "App pipeline: configure",
         "App owner approval", "Weekly plan", "Dev, Test, Prod", "Runs after Jira is healthy (step 8)"],
        ["T04", "C08", "Platform Bootstrap", "GitLab group/project for Jira IaC", "gitlab/platform/gitlab-config",
         ">= 1.6 (pinned)", "gitlabhq/gitlab@<ver>", "A13", "", "GitLab-managed TF state",
         "GitLab group access token from Vault", "gitlab_group, gitlab_project, gitlab_group_membership, protected branches",
         "Onboarding pipeline", "Platform approval", "Weekly plan", "All", "Runs once at onboarding"],
    ]
    build_table_sheet(wb, "6_Terraform", headers, rows, {"C": "TFPurpose", "P": "Environment"})


def build_network(wb: Workbook) -> None:
    headers = [
        ("Flow ID", 8, True), ("Source", 30, True), ("Destination", 30, True), ("Port(s)", 12, True), ("Protocol", 10, True),
        ("Direction", 16, True), ("Purpose", 38, True), ("Encrypted (TLS)?", 12, True), ("Environments", 14, True), ("Notes", 30, False),
    ]
    rows = [
        ["N01", "End users (corp network)", "Ingress jira.<env>.gdc.example.org", "443", "TCP", "Inbound", "Jira UI/API", "Yes", "All", ""],
        ["N02", "Ingress controller", "Jira pods", "8080", "TCP", "Internal (east-west)", "App traffic", "No (in-cluster)", "All", ""],
        ["N03", "Jira pods", "Jira pods", "40001, 40011", "TCP", "Internal (east-west)", "DC cluster cache replication", "No", "Prod", "Only when replicas > 1"],
        ["N04", "Jira pods (cluster egress)", "PostgreSQL VM", "5432", "TCP", "Outbound", "Database", "Yes", "All", ""],
        ["N05", "Jira pods", "SMTP relay", "587", "TCP", "Outbound", "Email notifications", "Yes", "All", ""],
        ["N06", "Jira pods", "AD / LDAP", "636", "TCP", "Outbound", "User directory", "Yes", "All", ""],
        ["N07", "K8s nodes", "JFrog Artifactory", "443", "TCP", "Outbound", "Image / chart pulls", "Yes", "All", ""],
        ["N08", "PostgreSQL VM", "JFrog Artifactory", "443", "TCP", "Outbound", "apt packages", "Yes", "All", ""],
        ["N09", "Ansible runner", "PostgreSQL VM", "22", "TCP", "Inbound", "Configuration", "Yes", "All", ""],
        ["N10", "Terraform runner", "Jira API", "443", "TCP", "Outbound", "Day-2 config (T03)", "Yes", "All", ""],
        ["N11", "PostgreSQL primary", "PostgreSQL standby", "5432", "TCP", "Internal (east-west)", "Streaming replication", "Yes", "Prod", ""],
        ["N12", "PostgreSQL VM", "Backup storage", "443", "TCP", "Outbound", "pgBackRest repository", "Yes", "All", ""],
    ]
    build_table_sheet(wb, "7_Network_Flows", headers, rows, {"F": "Direction", "H": "YesNo", "I": "Environment"})


def build_secrets(wb: Workbook) -> None:
    headers = [
        ("Secret ID", 9, True), ("Secret Name", 28, True), ("Used By (Component)", 16, True), ("Type", 22, True),
        ("Source of Truth", 26, True), ("Path / Reference (NO VALUES)", 38, True), ("Delivered Via", 34, True),
        ("Rotation", 16, True), ("Owner", 18, True), ("Notes", 30, False),
    ]
    rows = [
        ["S01", "jira-db-credentials", "C01, C02", "DB username/password", "HashiCorp Vault", "secret/apps/jira/<env>/db", "ExternalSecrets -> K8s Secret; Ansible lookup", "90 days", "Platform DBA", ""],
        ["S02", "postgres superuser", "C02", "DB password", "HashiCorp Vault", "secret/platform/pg/jira-<env>/postgres", "Ansible lookup", "90 days", "Platform DBA", ""],
        ["S03", "Jira license key", "C01", "License", "HashiCorp Vault", "secret/apps/jira/license", "Entered at setup / K8s Secret", "On renewal", "Collab Tools Team", ""],
        ["S04", "jira-tls", "C07", "TLS certificate", "Internal CA / cert-manager", "ClusterIssuer internal-ca", "cert-manager Certificate", "Auto (90 days)", "Platform Team", ""],
        ["S05", "jfrog-pull", "C04", "Registry pull token", "JFrog Access Token", "secret/platform/jfrog/pull/jira", "imagePullSecret", "30 days", "Platform Team", "Read-only, scoped to docker-virtual"],
        ["S06", "jira-admin-api-token", "C06", "API token", "HashiCorp Vault", "secret/apps/jira/<env>/tf-admin", "GitLab CI (Vault JWT auth)", "90 days", "Collab Tools Team", ""],
        ["S07", "ldap-bind", "C01", "Service account", "HashiCorp Vault", "secret/apps/jira/<env>/ldap", "Jira admin UI / K8s Secret", "180 days", "Identity Team", ""],
        ["S08", "gitlab-group-token", "C08", "API token", "HashiCorp Vault", "secret/platform/gitlab/tf", "GitLab CI (Vault JWT auth)", "90 days", "Platform Team", ""],
    ]
    build_table_sheet(wb, "8_Secrets_Config", headers, rows, {"E": "SecretStore"})


def build_sequence(wb: Workbook) -> None:
    headers = [
        ("Step", 6, True), ("Pipeline Stage", 22, True), ("Action", 50, True), ("Component / Stack", 16, True),
        ("Tool", 18, True), ("Depends On Step", 12, True), ("Success Criteria", 45, True), ("Rollback", 38, False),
    ]
    rows = [
        ["1", "ingest", "Pull/proxy all artifacts A01-A13 into JFrog, Xray scan, promote to virtual repos", "All", "JFrog CLI / Xray", "",
         "All artifacts present; no Critical/High CVEs without waiver", "Delete from local repo"],
        ["2", "validate", "helm template with JFrog overrides; assert no non-JFrog image refs", "C01", "Helm", "1",
         "0 images outside <jfrog> registry", "N/A"],
        ["3", "provision-infra", "Create namespace baseline", "T02", "Terraform", "1", "Namespace + quotas + pull secret exist", "terraform destroy (non-prod)"],
        ["4", "provision-infra", "Provision PostgreSQL VM(s)", "T01", "Terraform", "1", "VM reachable via SSH from Ansible runner", "terraform destroy (non-prod)"],
        ["5", "configure-vm", "Install & configure PostgreSQL 17, create jiradb/jirauser", "C02", "Ansible", "4",
         "psql connect as jirauser from K8s egress succeeds", "playbooks/postgresql-uninstall.yml"],
        ["6", "secrets", "Sync secrets S01, S03-S05 into namespace", "C04", "ExternalSecrets", "3, 5", "Secrets present in namespace", "N/A"],
        ["7", "deploy-app", "helm upgrade --install jira with env values", "C01", "Helm", "2, 6", "Pods Ready; /status returns RUNNING", "helm rollback"],
        ["8", "configure-app", "Apply Jira projects/groups", "T03", "Terraform", "7", "terraform plan shows no drift", "terraform apply previous state"],
        ["9", "verify", "Smoke tests: login, create issue, DB connectivity, email", "C01", "Test suite", "8", "All smoke tests pass", "Rollback steps 7-8"],
    ]
    build_table_sheet(wb, "9_Deploy_Sequence", headers, rows)


def build_checklist(wb: Workbook) -> None:
    headers = [
        ("#", 5, True), ("Category", 18, True), ("Check", 60, True), ("Owner", 18, True), ("Status", 14, True),
        ("Evidence / Link", 38, False), ("Notes", 40, False),
    ]
    rows = [
        ["1", "Vendor Support", "Vendor supported-platforms matrix confirms Jira version supports PostgreSQL 17", "Collab Tools Team", "Not Started", "", "Blocker if unsupported - fall back to supported PG major"],
        ["2", "Vendor Support", "Helm chart version compatible with Jira version and GDC K8s version", "Platform Team", "Not Started", "", ""],
        ["3", "Artifacts", "Every image rendered by 'helm template' is listed in 3_Artifact_Inventory", "Platform Team", "Not Started", "", ""],
        ["4", "Artifacts", "All artifacts pinned by version AND digest/checksum (no 'latest')", "App Team", "Not Started", "", ""],
        ["5", "Security", "Xray scan: no Critical/High CVEs, or approved waivers recorded", "Security", "Not Started", "", ""],
        ["6", "Security", "Licences reviewed and approved (incl. GPL collections, community providers)", "Legal / OSPO", "Not Started", "", ""],
        ["7", "Security", "Community Terraform providers reviewed (source, maintainer, signature)", "Security", "Not Started", "", "A12"],
        ["8", "Security", "No secret values stored in Git, values files or this workbook", "App Team", "Not Started", "", ""],
        ["9", "Infra", "RWX StorageClass available on target GDC cluster", "Platform Team", "Not Started", "", "C05"],
        ["10", "Infra", "Quota (CPU/RAM/storage/IP) available in GDC project", "Platform Team", "Not Started", "", ""],
        ["11", "Network", "All flows in 7_Network_Flows approved and firewall rules raised", "Network Team", "Not Started", "", ""],
        ["12", "Operations", "Backup + restore tested for PostgreSQL and Jira shared home", "Platform DBA", "Not Started", "", ""],
        ["13", "Operations", "Monitoring/alerts and runbooks in place", "App Team", "Not Started", "", ""],
        ["14", "Approval", "Ingestion approved by Platform & Security", "Platform Lead", "Not Started", "", ""],
    ]
    build_table_sheet(wb, "10_Readiness_Checklist", headers, rows, {"E": "Status"})


def build_lookups(wb: Workbook) -> None:
    ws = wb.create_sheet("Lookups")
    for col, (name, values) in enumerate(LOOKUPS.items(), start=1):
        ws.cell(row=1, column=col, value=name)
        for row, value in enumerate(values, start=2):
            ws.cell(row=row, column=col, value=value)
        ws.column_dimensions[get_column_letter(col)].width = 32
    style_header(ws, 1, len(LOOKUPS))


def main() -> None:
    wb = Workbook()
    build_instructions(wb)
    build_overview(wb)
    build_components(wb)
    build_artifacts(wb)
    build_k8s(wb)
    build_vm(wb)
    build_terraform(wb)
    build_network(wb)
    build_secrets(wb)
    build_sequence(wb)
    build_checklist(wb)
    build_lookups(wb)
    wb.save(OUTPUT)
    print(f"Created {OUTPUT}")


if __name__ == "__main__":
    main()
